"""Integration: PPD is generated alongside the žádost via /api/generate.

DATA_DIR is monkeypatched to a tmp dir so the test never touches the real
NAS evidence ledger / counter.
"""
import os

import app as appmod


def _payload(**over):
    base = {
        "mode": "prevod",
        "registracni_znacka": "1AB2345",
        "vin": "WBA3A5C51DF123456",
        "puvodni_jmeno": "JAN PRODÁVAJÍCÍ",
        "novy_jmeno": "PETR KUPUJÍCÍ",
        "ppd_castka": "1300",
        "ppd_prijato_od": "PETR KUPUJÍCÍ",
    }
    base.update(over)
    return base


def test_generate_includes_ppd(client, tmp_path, monkeypatch):
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    r = client.post("/api/generate", json=_payload())
    assert r.status_code == 200
    data = r.get_json()
    assert data["success"] is True
    assert data["ppd"].startswith("/download/ppd_")
    # file written + evidence row created under the tmp DATA_DIR
    fname = data["ppd"].split("/")[-1]
    assert os.path.exists(os.path.join(str(tmp_path), "output", fname))
    assert os.path.exists(os.path.join(str(tmp_path), "ppd_evidence.xlsx"))


def test_generate_pushes_to_evidence_by_default(client, tmp_path, monkeypatch):
    """The 'Zapsat úkon do evidence' box defaults on, so a normal generate fires
    the tracker push (with the žádost data)."""
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    import tracker_push
    calls = []
    monkeypatch.setattr(tracker_push, "push", lambda data, dd: calls.append(data))
    r = client.post("/api/generate", json=_payload())     # no evidence_log → default true
    assert r.get_json()["success"] is True
    assert len(calls) == 1


def test_generate_skips_evidence_when_unchecked(client, tmp_path, monkeypatch):
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    import tracker_push
    calls = []
    monkeypatch.setattr(tracker_push, "push", lambda data, dd: calls.append(data))
    r = client.post("/api/generate", json=_payload(evidence_log=False))
    assert r.get_json()["success"] is True
    assert calls == []                                    # box off → no push


def test_ppd_payer_is_uppercased(client, tmp_path, monkeypatch):
    """A hand-typed / autofilled payer is stored UPPERCASE on the receipt, like
    the rest of the form. The receipt ledger row reflects it."""
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    r = client.post("/api/generate", json=_payload(
        ppd_prijato_od="autodoprava novák s.r.o.", ppd_prijato_adresa="hlavní 5, brno"))
    assert r.get_json()["success"] is True
    import ppd as ppdmod
    log = ppdmod.read_ppd_log(str(tmp_path))
    assert log[0]["prijato_od"] == "AUTODOPRAVA NOVÁK S.R.O."


def test_ppd_print_page_is_a5(client, tmp_path, monkeypatch):
    """The generate response links an HTML print page whose @page rule pre-sets
    A5 paper in the browser's print dialog (the žádosti stay A4 PDFs)."""
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    r = client.post("/api/generate", json=_payload())
    data = r.get_json()
    assert data["ppd_print"].startswith("/ppd-print/")
    page = client.get(data["ppd_print"])
    assert page.status_code == 200
    html = page.get_data(as_text=True)
    assert "148mm 210mm" in html                    # @page = explicit A5 paper size
    # Misfeed tolerance: the frame is smaller than the printable area and
    # centered on the paper (~9-12 mm clearance), so a printer that feeds A5 a
    # few mm off can't cut the top border line (bug seen on real hardware).
    assert "width: 130mm" in html
    assert "height: 186mm" in html
    assert "no-store" in page.headers.get("Cache-Control", "")  # always fresh (no stale A4 page)
    assert "PŘÍJMOVÝ POKLADNÍ DOKLAD" in html
    assert "PETR KUPUJÍCÍ" in html
    assert "1AB2345" in html                        # SPZ on the receipt
    assert "window.print()" in html                 # auto-opens the dialog


def test_ppd_print_unknown_number_404(client, tmp_path, monkeypatch):
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    assert client.get("/ppd-print/999").status_code == 404


def test_zero_amount_skips_ppd(client, tmp_path, monkeypatch):
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    r = client.post("/api/generate", json=_payload(ppd_castka="0"))
    assert r.status_code == 200
    data = r.get_json()
    assert data["success"] is True
    assert "ppd" not in data            # opt-out
    assert data.get("zmeny")            # žádosti still produced


# ── víc vozidel na jednom dokladu ────────────────────────────────────────────
# Reálný případ: jeden PPD kryje víc aut najednou. Primární vozidlo je to ze
# žádosti (registracni_znacka); "+ Přidat vozidlo" v UI přidává další SPZ do
# ppd_extra_spz (čárkou oddělené), server je spojí do jednoho řádku.

def test_ppd_joins_extra_spz_from_the_expander(client, tmp_path, monkeypatch):
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    r = client.post("/api/generate", json=_payload(ppd_extra_spz="2CD6789, 3EF1234"))
    assert r.get_json()["success"] is True

    import openpyxl
    import ppd as ppdmod
    ws = openpyxl.load_workbook(ppdmod._evidence_path(str(tmp_path))).active
    row = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None][0]
    assert row[5] == "1AB2345, 2CD6789, 3EF1234"   # Vozidlo column


def test_ppd_ignores_blank_extra_entries(client, tmp_path, monkeypatch):
    """Trailing commas / empty rows left in the UI must not leave gaps like
    '1AB2345, , 3EF1234'."""
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    r = client.post("/api/generate", json=_payload(ppd_extra_spz="2CD6789, , "))
    assert r.get_json()["success"] is True
    import openpyxl
    import ppd as ppdmod
    ws = openpyxl.load_workbook(ppdmod._evidence_path(str(tmp_path))).active
    row = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None][0]
    assert row[5] == "1AB2345, 2CD6789"


def test_ppd_works_with_no_extra_spz(client, tmp_path, monkeypatch):
    """Default case (no expander touched) must be unchanged: just the one plate."""
    monkeypatch.setattr(appmod, "DATA_DIR", str(tmp_path))
    r = client.post("/api/generate", json=_payload())
    import openpyxl
    import ppd as ppdmod
    ws = openpyxl.load_workbook(ppdmod._evidence_path(str(tmp_path))).active
    row = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None][0]
    assert row[5] == "1AB2345"


def test_ppd_slovy_endpoint(client):
    r = client.get("/api/ppd-slovy?castka=1300")
    assert r.status_code == 200
    body = r.get_json()
    assert "tisíc" in body["slovy"].lower()


def test_ppd_slovy_endpoint_zero_and_bad_input(client):
    assert client.get("/api/ppd-slovy?castka=0").get_json()["slovy"]
    r = client.get("/api/ppd-slovy?castka=neplatne")
    assert r.status_code == 400
